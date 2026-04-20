"""
main.py — Customer Health Score Orchestrator

Pulls data from Stripe, Intercom, and HubSpot, computes a composite Health
Score for every active account, generates alerts, and exports the results to
a dated Excel workbook.

Usage:
    python main.py [--output-dir DIR] [--sync-hubspot] [--dry-run]
"""
from __future__ import annotations

import argparse
import logging
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from data_sources.hubspot_client import HubSpotDataExtractor
from data_sources.intercom_client import IntercomDataExtractor
from data_sources.stripe_client import StripeDataExtractor
from scoring.alerts import generate_alerts
from scoring.health_score import calculate_health_score

load_dotenv()

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("health_score")

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_CATEGORY_COLORS: dict[str, str] = {
    "Healthy":   "C6EFCE",
    "Stable":    "FFEB9C",
    "Watchlist": "FFCC99",
    "At Risk":   "FFC7CE",
}
_SEVERITY_COLORS: dict[str, str] = {
    "CRITICAL":    "FFC7CE",
    "HIGH":        "FFCC99",
    "MEDIUM":      "FFEB9C",
    "OPPORTUNITY": "C6EFCE",
    "STRATEGIC":   "BDD7EE",
}
_HEADER_BG = "1F4E79"
_HEADER_FG = "FFFFFF"

# Ordered columns for the "Health Scores" sheet.
_HS_COLS = [
    "customer_email", "customer_id", "status", "plan_name", "mrr",
    "tenure_months", "health_score", "category",
    "engagement_score", "financial_score", "support_score", "context_score",
    "failure_count_90d", "ghost_flag",
    "ticket_count_30d", "open_tickets_48h", "churn_mention",
    "in_cancel_pipeline", "in_unpaid_pipeline",
    "alert_count", "alert_ids",
]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Compute Customer Health Scores and export to Excel.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output"),
        metavar="DIR",
        help="Directory where the Excel file will be saved.",
    )
    p.add_argument(
        "--sync-hubspot",
        action="store_true",
        default=False,
        help="Push computed health_score values to HubSpot contact properties.",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        default=False,
        help="Run extraction and scoring but skip file writes and HubSpot sync.",
    )
    return p.parse_args()


# ---------------------------------------------------------------------------
# Extraction helpers
# ---------------------------------------------------------------------------

def _safe(fn, label: str, fallback):
    """Call fn(); on failure log a warning and return fallback."""
    try:
        result = fn()
        logger.info("  ✓ %s", label)
        return result
    except Exception as exc:
        logger.warning("  ⚠ %s failed — using empty fallback. %s: %s",
                       label, type(exc).__name__, exc)
        return fallback


def _extract_stripe(ext: StripeDataExtractor) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    logger.info("── Stripe ──────────────────────────────")
    subs = ext.get_all_active_subscriptions()          # required; no fallback
    logger.info("  ✓ Subscriptions (%d rows)", len(subs))

    customer_ids = subs["customer_id"].dropna().unique().tolist()

    failures = _safe(
        lambda: ext.get_payment_failures(customer_ids),
        f"Payment failures ({len(customer_ids)} customers)",
        pd.DataFrame(columns=["customer_id", "failure_count"]),
    )
    ghosts = _safe(
        ext.detect_ghost_invoices,
        "Ghost invoice detection",
        pd.DataFrame(columns=["customer_id", "ghost_flag"]),
    )
    return subs, failures, ghosts


def _extract_intercom(ext: IntercomDataExtractor) -> tuple[pd.DataFrame, pd.DataFrame]:
    logger.info("── Intercom ─────────────────────────────")
    tickets = _safe(
        lambda: ext.count_tickets_by_customer(days=30),
        "Ticket counts (30 days)",
        pd.DataFrame(columns=["email", "ticket_count", "open_tickets_48h"]),
    )
    churn = _safe(
        lambda: ext.detect_churn_mentions(days=14),
        "Churn-intent mentions (14 days)",
        pd.DataFrame(columns=["email", "churn_mention"]),
    )
    return tickets, churn


def _extract_hubspot(ext: HubSpotDataExtractor) -> tuple[set[str], set[str]]:
    logger.info("── HubSpot ──────────────────────────────")
    cancel_df = _safe(
        ext.get_deals_in_cancel_pipeline,
        "Cancel pipeline deals",
        pd.DataFrame(columns=["contact_email"]),
    )
    unpaid_df = _safe(
        ext.get_deals_in_unpaid_pipeline,
        "Unpaid pipeline deals",
        pd.DataFrame(columns=["contact_email"]),
    )
    return (
        set(cancel_df["contact_email"].dropna()),
        set(unpaid_df["contact_email"].dropna()),
    )


# ---------------------------------------------------------------------------
# Merge
# ---------------------------------------------------------------------------

def _merge(
    subs: pd.DataFrame,
    failures: pd.DataFrame,
    ghosts: pd.DataFrame,
    tickets: pd.DataFrame,
    churn: pd.DataFrame,
    cancel_emails: set[str],
    unpaid_emails: set[str],
) -> pd.DataFrame:
    """
    Join all sources into one row per customer (highest-MRR subscription wins
    when the same email has multiple subscriptions).
    """
    logger.info("── Merge ────────────────────────────────")
    df = subs.copy()

    # --- Stripe join-ons --------------------------------------------------
    df = df.merge(
        failures.rename(columns={"failure_count": "failure_count_90d"}),
        on="customer_id", how="left",
    )
    df["failure_count_90d"] = df["failure_count_90d"].fillna(0).astype(int)

    df = df.merge(ghosts, on="customer_id", how="left")
    df["ghost_flag"] = df["ghost_flag"].fillna(False).astype(bool)

    # --- Derived fields ---------------------------------------------------
    now = pd.Timestamp.now(tz="UTC")
    df["tenure_months"] = ((now - df["created"]).dt.days / 30.44).round(1)

    # Placeholders for fields not yet sourced from a live system.
    # Connect a product-analytics source here to replace these defaults.
    df["mrr_trend"]               = 0.0
    df["days_active_30d"]         = 0
    df["last_login_days"]         = 30
    df["features_adopted"]        = pd.NA
    df["features_available"]      = pd.NA
    df["has_billing_tickets"]     = False
    df["has_addons"]              = False
    df["end_client_growth"]       = pd.NA
    df["utilization_ratio"]       = pd.NA
    df["hs_above_85_months"]      = pd.NA
    df["mrr_declining_months"]    = pd.NA
    df["utilization_high_months"] = pd.NA

    # --- Intercom ---------------------------------------------------------
    if not tickets.empty:
        df = df.merge(
            tickets.rename(columns={
                "email":        "customer_email",
                "ticket_count": "ticket_count_30d",
            }),
            on="customer_email", how="left",
        )
    else:
        df["ticket_count_30d"] = 0
        df["open_tickets_48h"] = 0

    df["ticket_count_30d"] = df["ticket_count_30d"].fillna(0).astype(int)
    df["open_tickets_48h"] = df["open_tickets_48h"].fillna(0).astype(int)

    if not churn.empty:
        churn_emails = set(churn["email"].dropna())
        df["churn_mention"] = df["customer_email"].isin(churn_emails)
    else:
        df["churn_mention"] = False

    # --- HubSpot flags ----------------------------------------------------
    df["in_cancel_pipeline"] = df["customer_email"].isin(cancel_emails)
    df["in_unpaid_pipeline"]  = df["customer_email"].isin(unpaid_emails)

    # --- Deduplicate: one row per email -----------------------------------
    df = df.dropna(subset=["customer_email"])
    df = (
        df.sort_values("mrr", ascending=False, na_position="last")
          .drop_duplicates(subset=["customer_email"], keep="first")
          .reset_index(drop=True)
    )

    logger.info("  Merged dataset: %d unique accounts.", len(df))
    return df


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

def _score_and_alert(df: pd.DataFrame) -> pd.DataFrame:
    logger.info("── Scoring ──────────────────────────────")
    scores = df.apply(calculate_health_score, axis=1, result_type="expand")
    df = df.join(scores)
    logger.info("  Health scores computed.")

    df["alerts"]      = df.apply(generate_alerts, axis=1)
    df["alert_ids"]   = df["alerts"].apply(lambda a: ", ".join(x["alert_id"] for x in a))
    df["alert_count"] = df["alerts"].apply(len)
    logger.info("  Alerts generated.")
    return df


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def _explode_alerts(df: pd.DataFrame, prefix: str) -> pd.DataFrame:
    """One row per alert whose ID starts with prefix (e.g. 'R' or 'O')."""
    rows = []
    for _, row in df.iterrows():
        for alert in row.get("alerts", []):
            if alert["alert_id"].startswith(prefix):
                rows.append({
                    "customer_email": row.get("customer_email"),
                    "plan_name":      row.get("plan_name"),
                    "mrr":            row.get("mrr"),
                    "health_score":   row.get("health_score"),
                    "category":       row.get("category"),
                    **alert,
                })
    cols = ["customer_email", "plan_name", "mrr", "health_score",
            "category", "alert_id", "alert_name", "severity", "action"]
    return pd.DataFrame(rows, columns=cols) if rows else pd.DataFrame(columns=cols)


def _fmt_sheet(ws, df: pd.DataFrame, color_col: str | None = None) -> None:
    """Bold headers, freeze row 1, auto-width columns, optional cell coloring."""
    hdr_fill = PatternFill(start_color=_HEADER_BG, end_color=_HEADER_BG, fill_type="solid")
    hdr_font = Font(bold=True, color=_HEADER_FG)

    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"

    if color_col and color_col in df.columns:
        col_idx   = list(df.columns).index(color_col) + 1  # 1-based
        color_map = _CATEGORY_COLORS if color_col == "category" else _SEVERITY_COLORS
        for r, (_, data_row) in enumerate(df.iterrows(), start=2):
            val  = data_row.get(color_col, "")
            fill = PatternFill(
                start_color=color_map.get(str(val), "FFFFFF"),
                end_color=color_map.get(str(val), "FFFFFF"),
                fill_type="solid",
            )
            ws.cell(row=r, column=col_idx).fill = fill

    for col in ws.columns:
        letter  = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[letter].width = min(max_len + 3, 60)


def _build_summary(ws, df: pd.DataFrame) -> None:
    """Write the multi-section Resumen sheet directly via openpyxl."""
    bold12     = Font(bold=True, size=12)
    bold10     = Font(bold=True)
    hdr_fill   = PatternFill(start_color=_HEADER_BG, end_color=_HEADER_BG, fill_type="solid")
    hdr_font   = Font(bold=True, color=_HEADER_FG)

    total = len(df)
    r = 1

    # ── Distribution ──────────────────────────────────────────────────────
    ws.cell(r, 1, "Distribución por Categoría").font = bold12
    r += 1
    for c, h in enumerate(["Categoría", "Cuentas", "% del Total"], 1):
        ws.cell(r, c, h).fill = hdr_fill
        ws.cell(r, c, h).font = hdr_font
    r += 1
    for cat in ["Healthy", "Stable", "Watchlist", "At Risk"]:
        count = int((df["category"] == cat).sum())
        pct   = f"{count / total * 100:.1f}%" if total else "—"
        fill  = PatternFill(start_color=_CATEGORY_COLORS[cat],
                            end_color=_CATEGORY_COLORS[cat], fill_type="solid")
        for c, val in enumerate([cat, count, pct], 1):
            cell = ws.cell(r, c, val)
            cell.fill = fill
        r += 1

    r += 1  # blank row

    # ── Score statistics ──────────────────────────────────────────────────
    ws.cell(r, 1, "Estadísticas de Score").font = bold12
    r += 1
    stats = [
        ("Promedio",  round(float(df["health_score"].mean()),   1)),
        ("Mediana",   round(float(df["health_score"].median()), 1)),
        ("Mínimo",    round(float(df["health_score"].min()),    1)),
        ("Máximo",    round(float(df["health_score"].max()),    1)),
    ]
    for label, val in stats:
        ws.cell(r, 1, label).font = bold10
        ws.cell(r, 2, val)
        r += 1

    r += 1

    # ── Top 5 alerts ──────────────────────────────────────────────────────
    ws.cell(r, 1, "Top 5 Alertas más Frecuentes").font = bold12
    r += 1
    for c, h in enumerate(["Alert ID", "Nombre", "Frecuencia"], 1):
        ws.cell(r, c, h).fill = hdr_fill
        ws.cell(r, c, h).font = hdr_font
    r += 1

    all_alerts = [a for row_alerts in df["alerts"] for a in row_alerts]
    top5 = Counter((a["alert_id"], a["alert_name"]) for a in all_alerts).most_common(5)
    for (aid, aname), freq in top5:
        ws.cell(r, 1, aid)
        ws.cell(r, 2, aname)
        ws.cell(r, 3, freq)
        r += 1

    r += 1

    # ── Metadata ──────────────────────────────────────────────────────────
    critical = int(
        df["alerts"].apply(lambda a: any(x["severity"] == "CRITICAL" for x in a)).sum()
    )
    for label, val in [
        ("Generado el",                   str(date.today())),
        ("Total cuentas",                 total),
        ("Cuentas con alertas CRITICAL",  critical),
    ]:
        ws.cell(r, 1, label).font = bold10
        ws.cell(r, 2, val)
        r += 1

    for col, w in [("A", 36), ("B", 22), ("C", 14)]:
        ws.column_dimensions[col].width = w


def _write_excel(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    display_cols = [c for c in _HS_COLS if c in df.columns]
    hs_df   = df[display_cols].copy()
    risk_df = (
        hs_df[hs_df["category"] == "At Risk"]
        .sort_values("health_score")
        .reset_index(drop=True)
    )
    r_alerts = _explode_alerts(df, "R")
    o_alerts = _explode_alerts(df, "O")

    sheets = [
        ("Health Scores",       hs_df,    "category"),
        ("At Risk",             risk_df,  "category"),
        ("Alertas Riesgo",      r_alerts, "severity"),
        ("Alertas Oportunidad", o_alerts, "severity"),
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, sheet_df, _ in sheets:
            sheet_df.to_excel(writer, sheet_name=name, index=False)

        # Resumen sheet (manual openpyxl)
        ws_res = writer.book.create_sheet("Resumen")
        _build_summary(ws_res, df)

        # Format each data sheet
        for name, sheet_df, color_col in sheets:
            _fmt_sheet(writer.sheets[name], sheet_df, color_col)

    logger.info("  Saved: %s", path.resolve())


# ---------------------------------------------------------------------------
# Console summary
# ---------------------------------------------------------------------------

def _log_summary(df: pd.DataFrame) -> None:
    total    = len(df)
    avg      = df["health_score"].mean()
    critical = int(
        df["alerts"].apply(lambda a: any(x["severity"] == "CRITICAL" for x in a)).sum()
    )
    cat_counts = df["category"].value_counts()

    sep = "─" * 52
    logger.info(sep)
    logger.info("HEALTH SCORE REPORT  —  %s", date.today())
    logger.info(sep)
    logger.info("Cuentas procesadas   : %d", total)
    logger.info("Score promedio       : %.1f", avg)
    logger.info("")
    logger.info("Distribución:")
    for cat in ["Healthy", "Stable", "Watchlist", "At Risk"]:
        n   = int(cat_counts.get(cat, 0))
        bar = "█" * min(n, 30)
        logger.info("  %-12s %4d  %s", cat, n, bar)
    logger.info("")
    logger.info("Alertas CRITICAL activas: %d", critical)

    all_alerts = [a for row_a in df["alerts"] for a in row_a]
    if all_alerts:
        top3 = Counter(a["alert_id"] for a in all_alerts).most_common(3)
        logger.info("Top alertas          : %s",
                    "  ".join(f"{aid}×{n}" for aid, n in top3))
    logger.info(sep)


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------

def main() -> None:
    args = _parse_args()

    try:
        logger.info("Initialising extractors…")
        stripe_ext   = StripeDataExtractor()
        intercom_ext = IntercomDataExtractor()
        hubspot_ext  = HubSpotDataExtractor()

        subs, failures, ghosts       = _extract_stripe(stripe_ext)
        tickets, churn               = _extract_intercom(intercom_ext)
        cancel_emails, unpaid_emails = _extract_hubspot(hubspot_ext)

        df = _merge(subs, failures, ghosts,
                    tickets, churn,
                    cancel_emails, unpaid_emails)

        if df.empty:
            logger.warning("No accounts to process — exiting.")
            sys.exit(0)

        df = _score_and_alert(df)
        _log_summary(df)

        if args.dry_run:
            logger.info("Dry-run active — skipping file write and HubSpot sync.")
            return

        if args.sync_hubspot:
            logger.info("── HubSpot sync ─────────────────────────")
            hubspot_ext.sync_health_score(
                df[["customer_email", "health_score"]]
                  .rename(columns={"customer_email": "email"})
            )

        logger.info("── Excel export ─────────────────────────")
        output_path = args.output_dir / f"health_score_{date.today()}.xlsx"
        _write_excel(df, output_path)

    except KeyboardInterrupt:
        logger.info("Interrupted by user.")
        sys.exit(0)

    except EnvironmentError as exc:
        logger.error("Configuration error: %s", exc)
        sys.exit(1)

    except Exception as exc:
        logger.exception("Fatal error: %s", exc)
        sys.exit(1)


if __name__ == "__main__":
    main()
