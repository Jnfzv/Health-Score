"""
output/exporter.py — Professional Excel export for Customer Health Scores.

Public API:
    export_to_excel(scores_df, alerts_df, output_path)

scores_df expected columns (any subset; missing columns are left blank):
    customer_email / email, plan_name / plan, status,
    health_score, category,
    engagement_score / engagement, financial_score / financial,
    support_score / support, context_score / context,
    alert_ids / active_alerts, last_login_days, tenure_months

alerts_df expected columns (exploded, one row per alert):
    customer_email / email, alert_id, alert_name, severity, action
    (plus any extra columns — only the above are used)
"""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Colour constants
# ---------------------------------------------------------------------------

_HEADER_BG   = "2F4F6F"   # dark slate blue
_HEADER_FG   = "FFFFFF"

_SCORE_COLORS = {
    "green":  "C6EFCE",   # Healthy   >= 75
    "yellow": "FFEB9C",   # Stable    55-74
    "orange": "FFCC99",   # Watchlist 35-54
    "red":    "FFC7CE",   # At Risk   < 35
}
_CATEGORY_COLORS = {
    "Healthy":   _SCORE_COLORS["green"],
    "Stable":    _SCORE_COLORS["yellow"],
    "Watchlist": _SCORE_COLORS["orange"],
    "At Risk":   _SCORE_COLORS["red"],
}
_SEVERITY_COLORS = {
    "CRITICAL":    _SCORE_COLORS["red"],
    "HIGH":        _SCORE_COLORS["orange"],
    "MEDIUM":      _SCORE_COLORS["yellow"],
    "OPPORTUNITY": _SCORE_COLORS["green"],
    "STRATEGIC":   "BDD7EE",
}


# ---------------------------------------------------------------------------
# Column specification for "Health Scores" sheet
# Each entry: (source_column_candidates, display_header, min_width)
# ---------------------------------------------------------------------------

_HS_SPEC: list[tuple[list[str], str, int]] = [
    (["customer_email", "email"],           "Email",           34),
    (["plan_name",      "plan"],            "Plan",            12),
    (["status"],                            "Status",          12),
    (["health_score"],                      "Health Score",    13),
    (["category"],                          "Category",        12),
    (["engagement_score", "engagement"],    "Engagement",      12),
    (["financial_score",  "financial"],     "Financial",       12),
    (["support_score",    "support"],       "Support",         12),
    (["context_score",    "context"],       "Context",         12),
    (["alert_ids", "active_alerts"],        "Active Alerts",   28),
    (["last_login_days"],                   "Last Login (d)",  15),
    (["tenure_months"],                     "Tenure (mo)",     13),
]


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _pick(df: pd.DataFrame, candidates: list[str]):
    """Return the Series for the first candidate column found, or None."""
    for c in candidates:
        if c in df.columns:
            return df[c]
    return None


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _header_font() -> Font:
    return Font(bold=True, color=_HEADER_FG, name="Calibri", size=10)


def _write_row(ws, row_idx: int, values: list, font=None, fill=None, center=False) -> None:
    align = Alignment(horizontal="center") if center else Alignment(horizontal="left")
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = align


def _section_title(ws, row_idx: int, text: str) -> None:
    cell = ws.cell(row=row_idx, column=1, value=text)
    cell.font = Font(bold=True, size=12, color="1F4E79", name="Calibri")


def _set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _auto_widen(ws, df: pd.DataFrame, min_widths: list[int]) -> None:
    """Set column widths to max(min_width, actual content width + padding)."""
    for i, (col, min_w) in enumerate(zip(df.columns, min_widths), start=1):
        data_max = df[col].astype(str).str.len().max() if not df.empty else 0
        header_len = len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = max(
            min_w, data_max + 3, header_len + 3
        )


# ---------------------------------------------------------------------------
# "Health Scores" sheet
# ---------------------------------------------------------------------------

def _build_hs_df(scores_df: pd.DataFrame) -> pd.DataFrame:
    """Extract and rename columns for the Health Scores sheet."""
    out = pd.DataFrame()
    for candidates, header, _ in _HS_SPEC:
        series = _pick(scores_df, candidates)
        out[header] = series if series is not None else pd.NA
    return out


def _apply_score_cf(ws, col_letter: str, last_row: int) -> None:
    """
    Add Excel-native conditional formatting on the Health Score column.
    Rules are mutually exclusive (between / greaterThanOrEqual / lessThan)
    so no stopIfTrue priority juggling is needed.
    """
    cell_range = f"{col_letter}2:{col_letter}{last_row}"

    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="greaterThanOrEqual", formula=["75"],
            fill=_fill(_SCORE_COLORS["green"]),
        ),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between", formula=["55", "74.99"],
            fill=_fill(_SCORE_COLORS["yellow"]),
        ),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between", formula=["35", "54.99"],
            fill=_fill(_SCORE_COLORS["orange"]),
        ),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="lessThan", formula=["35"],
            fill=_fill(_SCORE_COLORS["red"]),
        ),
    )


def _write_hs_sheet(ws, scores_df: pd.DataFrame) -> None:
    df = _build_hs_df(scores_df)

    headers    = [h for _, h, _ in _HS_SPEC]
    min_widths = [w for _, _, w in _HS_SPEC]
    hdr_fill   = _fill(_HEADER_BG)
    hdr_font   = _header_font()
    center     = Alignment(horizontal="center")

    # ── Header row ────────────────────────────────────────────────────────
    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(df) + 1, 1)}"
    ws.row_dimensions[1].height = 18

    # ── Data rows ─────────────────────────────────────────────────────────
    cat_col_idx = headers.index("Category") + 1   # 1-based
    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for c_idx, col in enumerate(df.columns, start=1):
            val  = row[col]
            cell = ws.cell(row=r_idx, column=c_idx,
                           value=None if pd.isna(val) else val)
            cell.alignment = Alignment(horizontal="center")

        # Category cell: coloured background
        cat_val  = row.get("Category", "")
        cat_fill = _fill(_CATEGORY_COLORS.get(str(cat_val), "FFFFFF"))
        ws.cell(row=r_idx, column=cat_col_idx).fill = cat_fill

    # ── Score conditional formatting ─────────────────────────────────────
    hs_col_idx = headers.index("Health Score") + 1
    if len(df) > 0:
        _apply_score_cf(
            ws,
            get_column_letter(hs_col_idx),
            len(df) + 1,
        )

    # ── Column widths ─────────────────────────────────────────────────────
    _auto_widen(ws, df, min_widths)


# ---------------------------------------------------------------------------
# "Resumen" sheet
# ---------------------------------------------------------------------------

def _write_resumen_sheet(ws, scores_df: pd.DataFrame, alerts_df: pd.DataFrame) -> None:
    hdr_fill  = _fill(_HEADER_BG)
    hdr_font  = _header_font()
    bold      = Font(bold=True, name="Calibri", size=10)
    center    = Alignment(horizontal="center")

    total     = len(scores_df)
    r         = 1   # current row pointer

    # ── 1. Category distribution ──────────────────────────────────────────
    _section_title(ws, r, "Distribución por Categoría")
    r += 1

    for c_idx, h in enumerate(["Categoría", "Cuentas", "% del Total"], 1):
        cell = ws.cell(row=r, column=c_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center
    r += 1

    cat_col = _pick(scores_df, ["category"])
    for cat in ["Healthy", "Stable", "Watchlist", "At Risk"]:
        count = int((cat_col == cat).sum()) if cat_col is not None else 0
        pct   = f"{count / total * 100:.1f}%" if total else "—"
        row_fill = _fill(_CATEGORY_COLORS[cat])
        for c_idx, val in enumerate([cat, count, pct], 1):
            cell = ws.cell(row=r, column=c_idx, value=val)
            cell.fill      = row_fill
            cell.alignment = center
        r += 1

    r += 1  # blank

    # ── 2. Score statistics ───────────────────────────────────────────────
    _section_title(ws, r, "Score Promedio Global")
    r += 1

    hs_col = _pick(scores_df, ["health_score"])
    if hs_col is not None and hs_col.notna().any():
        stats = [
            ("Promedio",  round(float(hs_col.mean()),   1)),
            ("Mediana",   round(float(hs_col.median()), 1)),
            ("Mínimo",    round(float(hs_col.min()),    1)),
            ("Máximo",    round(float(hs_col.max()),    1)),
        ]
        for label, val in stats:
            ws.cell(row=r, column=1, value=label).font = bold
            ws.cell(row=r, column=2, value=val).alignment = center
            r += 1
    else:
        ws.cell(row=r, column=1, value="Sin datos").font = bold
        r += 1

    r += 1

    # ── 3. Top 10 critical accounts ───────────────────────────────────────
    _section_title(ws, r, "Top 10 Cuentas Más Críticas")
    r += 1

    top_headers = ["Email", "Plan", "Health Score", "Category", "Active Alerts"]
    for c_idx, h in enumerate(top_headers, 1):
        cell = ws.cell(row=r, column=c_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center
    r += 1

    if hs_col is not None:
        top10 = (
            scores_df[hs_col.notna()]
            .nsmallest(10, hs_col.name)
            .reset_index(drop=True)
        )
        for _, row_data in top10.iterrows():
            email   = _pick_val(row_data, ["customer_email", "email"])
            plan    = _pick_val(row_data, ["plan_name", "plan"])
            score   = _pick_val(row_data, ["health_score"])
            cat     = _pick_val(row_data, ["category"])
            alerts  = _pick_val(row_data, ["alert_ids", "active_alerts", "alert_count"])
            row_fill = _fill(_CATEGORY_COLORS.get(str(cat), "FFFFFF"))
            for c_idx, val in enumerate([email, plan, score, cat, alerts], 1):
                cell = ws.cell(row=r, column=c_idx, value=val)
                cell.alignment = center
                if c_idx == 4:   # Category column gets color fill
                    cell.fill = row_fill
            r += 1

    r += 1

    # ── 4. Alert frequency ────────────────────────────────────────────────
    _section_title(ws, r, "Conteo de Alertas por Tipo")
    r += 1

    for c_idx, h in enumerate(["Alert ID", "Nombre", "Severidad", "Frecuencia"], 1):
        cell = ws.cell(row=r, column=c_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center
    r += 1

    if alerts_df is not None and not alerts_df.empty:
        id_col  = _pick(alerts_df, ["alert_id"])
        nm_col  = _pick(alerts_df, ["alert_name"])
        sev_col = _pick(alerts_df, ["severity"])

        if id_col is not None:
            freq = (
                alerts_df
                .assign(
                    _id  = id_col,
                    _nm  = nm_col  if nm_col  is not None else "",
                    _sev = sev_col if sev_col is not None else "",
                )
                .groupby(["_id", "_nm", "_sev"], dropna=False)
                .size()
                .reset_index(name="count")
                .sort_values("count", ascending=False)
            )
            for _, row_data in freq.iterrows():
                sev   = row_data["_sev"]
                row_fill = _fill(_SEVERITY_COLORS.get(str(sev), "FFFFFF"))
                for c_idx, val in enumerate(
                    [row_data["_id"], row_data["_nm"], sev, row_data["count"]], 1
                ):
                    cell = ws.cell(row=r, column=c_idx, value=val)
                    cell.alignment = center
                    cell.fill = row_fill
                r += 1
    else:
        ws.cell(row=r, column=1, value="Sin alertas activas").font = bold
        r += 1

    # ── Column widths ────────────────────────────────────────────────────
    for col_letter, width in [("A", 36), ("B", 16), ("C", 15), ("D", 13), ("E", 28)]:
        ws.column_dimensions[col_letter].width = width


def _pick_val(row, candidates: list[str]):
    """Get the first available value from a pandas Series (row) by candidate key."""
    for c in candidates:
        if c in row.index:
            val = row[c]
            return None if pd.isna(val) else val
    return None


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_to_excel(
    scores_df: pd.DataFrame,
    alerts_df: pd.DataFrame | None,
    output_path: str | Path,
) -> Path:
    """
    Write a formatted Excel workbook with two sheets:

    - "Health Scores": full customer table with conditional score formatting,
      auto-filtered headers, and category colour coding.
    - "Resumen": category distribution, score statistics, top-10 critical
      accounts, and alert frequency table.

    Parameters
    ----------
    scores_df:
        One row per customer, containing at minimum health_score and category.
    alerts_df:
        Exploded alerts (one row per active alert per customer). Pass None or
        an empty DataFrame to leave the alert section blank.
    output_path:
        Destination .xlsx path. Parent directories are created if absent.

    Returns the resolved Path of the written file.
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if alerts_df is None:
        alerts_df = pd.DataFrame(
            columns=["customer_email", "alert_id", "alert_name", "severity", "action"]
        )

    wb = Workbook()

    # Remove the default blank sheet created by openpyxl.
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    ws_hs = wb.create_sheet("Health Scores")
    _write_hs_sheet(ws_hs, scores_df)
    logger.debug("Sheet 'Health Scores' written (%d rows).", len(scores_df))

    ws_res = wb.create_sheet("Resumen")
    _write_resumen_sheet(ws_res, scores_df, alerts_df)
    logger.debug("Sheet 'Resumen' written.")

    wb.save(path)
    logger.info("Excel exported → %s", path.resolve())
    return path.resolve()
